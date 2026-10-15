# -*- coding: utf-8 -*-
"""
P&L-агент ТОО «Caspian Sea Support» (CSS).
Логика как в Градексе: читает выгрузку 1С ОСВ (Анализ счёта / ОСВ),
берёт только 6xxx (доходы) и 7xxx (расходы) и строит P&L по проектам.

Структура ОСВ:
  6010  доход: вид услуги (ур.2) -> проект (ур.3)         -> Кредит (кол.6)
  7011  затраты: статья (ур.3) -> проект (ур.4)           -> Дебет  (кол.5)
  7210/7211/7212 административные: статья (ур.3)           -> Дебет  (кол.5)
  6110 депозит, 6250 курс.доход, 6290 прочий доход        -> низ
  7300 финансирование, 7430 курс.расход, 7480 прочие      -> низ
  5610 нераспределённая прибыль (контроль)
"""
import re, openpyxl
from collections import defaultdict, OrderedDict

THRESHOLD = 300_000  # порог «новой» статьи/проекта (для флагов)


def norm(s):
    return re.sub(r'\s+', ' ', str(s or '').strip().lower()).replace('ё', 'е')


def _num(x):
    return float(x) if isinstance(x, (int, float)) else 0.0


# ---- канонизация имени проекта (разные написания -> одно) ----
def canon_project(p):
    n = norm(p)
    n = n.replace('проект ', '').replace('ncoc tar:', 'ncoc:').replace('ncoc:', '').replace('ncoc ', '')
    n = n.strip(' :')
    # вид судна/проекта
    if 'dredging' in n: return 'Dredging'
    if 'srp' in n: return 'SRP'
    if 'tulpar' in n: return 'TULPAR'
    if 'ibeev' in n: return 'IBEEV'
    if 'lqb' in n: return 'LQB'
    if 'ips' in n: return 'IPS'
    if 'jan de nul' in n and 'aktau' in n: return 'Jan De Nul (Aktau Port)'
    if 'jan de nul' in n or 'янденул' in n: return 'Jan De Nul'
    if 'ледозащит' in n or 'ui60102' in n: return 'Ледозащитные конструкции'
    if 'barges' in n or 'барж' in n or 'flat top' in n:
        if 'top 14' in n or 'top14' in n: return 'Flat Top Barges (TOP 14)'
        if 'сом 5' in n or 'com 5' in n or 'сом5' in n: return 'Flat Top Barges (СОМ 5)'
        return 'Flat Top Barges'
    if 'survey' in n or 'caspian maria' in n or 'caspian spider' in n or 'coastal bigfoot' in n \
       or 'veritas' in n or 'ocean tiara' in n or 'long term' in n:
        for key, lbl in [('long term', 'Survey Vessels (Long Term)'),
                          ('caspian maria', 'Survey (Caspian Maria)'),
                          ('caspian spider', 'Survey (Caspian Spider)'),
                          ('coastal bigfoot', 'Survey (Coastal Bigfoot)'),
                          ('veritas', 'Survey (Veritas Catro III)'),
                          ('ocean tiara', 'Survey (Ocean Tiara)')]:
            if key in n: return lbl
        return 'Survey Vessels'
    if 'usdt' in n or 'meric' in n:
        if 'eva' in n: return 'USDT Eva'
        if 'maria' in n: return 'USDT Maria'
        if 'meric' in n: return 'USDT Meric'
        return 'USDT'
    if 'основн' in n or 'номенклатур' in n: return 'Прочие (без проекта)'
    return p.strip() if p else 'Прочие (без проекта)'


def parse_1c(path):
    """Разбирает ОСВ. Возвращает структуру для построения P&L."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in range(1, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if name is None:
            continue
        name = str(name).strip()
        if name in ('<...>', ''):
            continue
        od = ws.row_dimensions[r].outline_level if r in ws.row_dimensions else 0
        deb = _num(ws.cell(row=r, column=5).value)
        cred = _num(ws.cell(row=r, column=6).value)
        rows.append((r, od, name, deb, cred))

    R = {
        'income': defaultdict(float),          # project -> выручка (6010)
        'income_services': defaultdict(lambda: defaultdict(float)),  # project -> {service: amt}
        'cost': defaultdict(lambda: defaultdict(float)),  # project -> {article: amt} (7011)
        'admin': defaultdict(float),           # article -> amt (7211/7212)
        'deposit': 0.0,                        # 6110
        'fx_income': 0.0,                      # 6250
        'other_income': 0.0,                   # 6290
        'financing': 0.0,                      # 7300
        'fx_expense': 0.0,                      # 7430
        'other_expense': 0.0,                  # 7480
        'ctrl_5610': 0.0,
        'tot_income_6': 0.0,
        'tot_expense_7': 0.0,
        'cost_target': 0.0,    # 7010 (ур.1) — для сверки прямых затрат
        'admin_target': 0.0,   # 7210 (ур.1) — для сверки админа
        'proj_raw': defaultdict(set),  # блок -> {сырые имена проектов из ОСВ}
    }

    acct = None          # текущий счёт верхнего уровня
    cur_service = None    # текущая статья/вид услуги (ур.2 для дохода, ур.3 для затрат)
    stack = {}

    for i, (r, od, name, deb, cred) in enumerate(rows):
        stack[od] = name
        for k in [k for k in stack if k > od]:
            del stack[k]
        nxt = rows[i + 1][1] if i + 1 < len(rows) else 0
        is_leaf = nxt <= od

        m = re.match(r'^(\d{4}),', name)
        if m:
            acct = m.group(1)
            # обороты по счёту верхнего уровня (для контроля сходимости)
            if od == 0:
                if acct.startswith('6'):
                    R['tot_income_6'] += cred
                elif acct.startswith('7') and acct < '7700':
                    R['tot_expense_7'] += deb
            if acct == '5610':
                R['ctrl_5610'] += cred - deb
            if od == 1:
                if acct == '7010':
                    R['cost_target'] += deb
                elif acct == '7210':
                    R['admin_target'] += deb
                # «нижние»/финансовые счета: сумма на строке субсчёта (ур.1)
                elif acct == '6110':
                    R['deposit'] += cred
                elif acct == '6250':
                    R['fx_income'] += cred
                elif acct == '6290':
                    R['other_income'] += cred
                elif acct in ('7310', '7320'):
                    R['financing'] += deb
                elif acct == '7430':
                    R['fx_expense'] += deb
                elif acct == '7480':
                    R['other_expense'] += deb
            # 7012 (и т.п.): сумма висит на строке субсчёта без разбивки
            if acct in ('7011', '7012') and od == 2 and is_leaf and deb:
                R['cost']['Прочие (без проекта)']['Себестоимость (прочее)'] += deb
            continue

        a = acct or ''
        # ---------- ДОХОД ОТ РЕАЛИЗАЦИИ 6010 (проект = ур.3) ----------
        if a == '6010':
            if od == 2:
                cur_service = name
            if od == 3 and cred:
                proj = canon_project(name)
                R['income'][proj] += cred
                R['income_services'][proj][cur_service or name] += cred
                R['proj_raw'][proj].add(name.strip())
            continue
        # ---------- ПРЯМЫЕ ЗАТРАТЫ 7011/7012 (проект = ур.4) ----------
        if a in ('7011', '7012'):
            if od == 3:
                cur_service = name
            if od == 4 and deb:
                proj = canon_project(name)
                R['cost'][proj][cur_service or name] += deb
                R['proj_raw'][proj].add(name.strip())
            elif od == 3 and is_leaf and deb:
                R['cost']['Прочие (без проекта)'][name] += deb
            continue
        # ---------- АДМИН 7210/7211/7212 ----------
        if a in ('7210', '7211', '7212'):
            if is_leaf and deb and od >= 3 and name != '<...>':
                R['admin'][name.strip()] += deb
            continue

    # ---- сверка остатков по целям (чтобы блоки сходились точно) ----
    cost_sum = sum(sum(a.values()) for a in R['cost'].values())
    resid = R['cost_target'] - cost_sum
    if abs(resid) > 1:
        R['cost']['Прочие (без проекта)']['Нераспознанные затраты'] += resid
    admin_sum = sum(R['admin'].values())
    resid_a = R['admin_target'] - admin_sum
    if abs(resid_a) > 1:
        R['admin']['Прочие (нераспознанное)'] = R['admin'].get('Прочие (нераспознанное)', 0) + resid_a

    return R


# =====================================================================
#  ГРУППИРОВКА АДМИНА (налоги с ЗП в одну строку, как в шаблоне)
# =====================================================================
_TAX_KEYS = ('обязательные пенсионные взносы', 'отчисления осмс',
             'социальные отчисления', 'социальный налог')

def group_admin(admin):
    """Возвращает OrderedDict статья->сумма, схлопывая налоги с ЗП в одну строку."""
    out = OrderedDict()
    tax = 0.0
    for art, amt in sorted(admin.items(), key=lambda x: -x[1]):
        if any(k in norm(art) for k in _TAX_KEYS):
            tax += amt
        else:
            out[art] = out.get(art, 0.0) + amt
    if tax:
        out['Налоги и обяз. платежи по ЗП'] = tax
    return out


# =====================================================================
#  ПОСТРОЕНИЕ P&L-КНИГИ
# =====================================================================
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

_THIN = Side(style='thin', color='D9D9D9')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_F_NAME = 'Arial'

def _style_num(c):
    c.number_format = '#,##0;(#,##0);-'
    c.alignment = Alignment(horizontal='right')
    c.font = Font(name=_F_NAME, size=10)
    c.border = _BORDER

def build_workbook(periods, out_path, company='ТОО «Caspian Sea Support»'):
    """periods: список (имя_колонки, R_parsed). Строит единый лист P&L."""
    # --- объединённый список проектов (по убыванию выручки по всем периодам) ---
    proj_rev = defaultdict(float)
    proj_set = set()
    for _, R in periods:
        for p, v in R['income'].items():
            proj_rev[p] += v; proj_set.add(p)
        for p in R['cost']:
            proj_set.add(p)
    # порядок: по суммарной выручке убыв., проекты без выручки — в конец
    projects = sorted(proj_set, key=lambda p: (-proj_rev.get(p, 0), p))

    # --- объединённый список статей затрат по каждому проекту ---
    proj_articles = {}
    for p in projects:
        arts = defaultdict(float)
        for _, R in periods:
            for a, v in R['cost'].get(p, {}).items():
                arts[a] += v
        proj_articles[p] = [a for a, _ in sorted(arts.items(), key=lambda x: -x[1])]

    # --- объединённый админ ---
    admin_all = defaultdict(float)
    for _, R in periods:
        for a, v in group_admin(R['admin']).items():
            admin_all[a] += v
    admin_lines = [a for a, _ in sorted(admin_all.items(), key=lambda x: -x[1])]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'P&L 2026'
    ncol = len(periods)
    data_cols = [get_column_letter(3 + i) for i in range(ncol)]   # C, D, ...
    tot_col = get_column_letter(3 + ncol)                         # ИТОГО

    # значения по (ключ_строки -> {период_idx -> сумма})
    def val(R, key):
        return R.get(key, 0.0)

    r = 1
    # заголовок
    ws.cell(row=r, column=1, value=company).font = Font(name=_F_NAME, size=12, bold=True)
    for i, (nm, _) in enumerate(periods):
        c = ws.cell(row=r, column=3 + i, value=nm)
        c.font = Font(name=_F_NAME, size=10, bold=True); c.alignment = Alignment(horizontal='center')
    c = ws.cell(row=r, column=3 + ncol, value='ИТОГО')
    c.font = Font(name=_F_NAME, size=10, bold=True); c.alignment = Alignment(horizontal='center')
    r += 1

    def section(title):
        nonlocal r
        cell = ws.cell(row=r, column=2, value=title)
        cell.font = Font(name=_F_NAME, size=10, bold=True, color='FFFFFF')
        for cc in range(2, 4 + ncol):
            ws.cell(row=r, column=cc).fill = PatternFill('solid', fgColor='2F5496')
            ws.cell(row=r, column=cc).border = _BORDER
        r += 1
        return r - 1

    def leaf(label, getter, indent=1):
        """строка-лист: label, значения по периодам (getter(R)), ИТОГО=SUM."""
        nonlocal r
        lc = ws.cell(row=r, column=2, value=('   ' * indent) + label)
        lc.font = Font(name=_F_NAME, size=10); lc.border = _BORDER
        for i, (_, R) in enumerate(periods):
            c = ws.cell(row=r, column=3 + i, value=round(getter(R), 2))
            _style_num(c)
        tc = ws.cell(row=r, column=3 + ncol,
                     value=f'=SUM({data_cols[0]}{r}:{data_cols[-1]}{r})' if ncol > 1 else f'={data_cols[0]}{r}')
        _style_num(tc); tc.font = Font(name=_F_NAME, size=10, bold=True)
        r += 1
        return r - 1

    def total(label, rowrefs=None, formula=None, bold=True, fill='D9E1F2'):
        """итоговая строка: по каждому периоду либо SUM(указанных строк), либо своя формула."""
        nonlocal r
        lc = ws.cell(row=r, column=2, value=label)
        lc.font = Font(name=_F_NAME, size=10, bold=bold); lc.border = _BORDER
        if fill:
            for cc in range(2, 4 + ncol):
                ws.cell(row=r, column=cc).fill = PatternFill('solid', fgColor=fill)
        for col in data_cols + [tot_col]:
            if formula:
                f = formula(col)
            elif rowrefs:
                f = '=' + '+'.join(f'{col}{rr}' for rr in rowrefs) if rowrefs else '=0'
            else:
                f = '=0'
            c = ws.cell(row=r, column=openpyxl.utils.column_index_from_string(col), value=f)
            _style_num(c); c.font = Font(name=_F_NAME, size=10, bold=bold)
        r += 1
        return r - 1

    # ---------------- ДОХОД ОТ РЕАЛИЗАЦИИ ----------------
    section('ДОХОД ОТ РЕАЛИЗАЦИИ (по проектам)')
    inc_rows = []
    for p in projects:
        if proj_rev.get(p, 0) == 0 and all(R['income'].get(p, 0) == 0 for _, R in periods):
            continue
        inc_rows.append(leaf(p, lambda R, p=p: R['income'].get(p, 0.0)))
    row_income_total = total('ИТОГО ДОХОД ОТ РЕАЛИЗАЦИИ', rowrefs=inc_rows, fill='BDD7EE')
    r += 1

    # ---------------- ПРЯМЫЕ ЗАТРАТЫ ПО ПРОЕКТАМ ----------------
    section('ПРЯМЫЕ ЗАТРАТЫ ПО ПРОЕКТАМ')
    proj_subtotal_rows = []
    for p in projects:
        arts = proj_articles.get(p, [])
        if not arts:
            continue
        # подзаголовок проекта
        hc = ws.cell(row=r, column=2, value=p)
        hc.font = Font(name=_F_NAME, size=10, bold=True, italic=True); hc.border = _BORDER
        for cc in range(2, 4 + ncol):
            ws.cell(row=r, column=cc).fill = PatternFill('solid', fgColor='EDEDED')
        r += 1
        art_rows = []
        for a in arts:
            art_rows.append(leaf(a, lambda R, p=p, a=a: R['cost'].get(p, {}).get(a, 0.0), indent=2))
        sub = total(f'Итого {p}', rowrefs=art_rows, bold=True, fill='F2F2F2')
        proj_subtotal_rows.append(sub)
    row_direct_total = total('ИТОГО ПРЯМЫЕ ЗАТРАТЫ', rowrefs=proj_subtotal_rows, fill='BDD7EE')
    # валовая прибыль
    row_gross = total('ВАЛОВАЯ ПРИБЫЛЬ',
                      formula=lambda col: f'={col}{row_income_total}-{col}{row_direct_total}',
                      fill='FFF2CC')
    r += 1

    # ---------------- АДМИНИСТРАТИВНЫЕ РАСХОДЫ ----------------
    section('АДМИНИСТРАТИВНЫЕ РАСХОДЫ')
    adm_rows = []
    for a in admin_lines:
        gdict = {nm: group_admin(R['admin']) for nm, R in periods}
        adm_rows.append(leaf(a, lambda R, a=a: group_admin(R['admin']).get(a, 0.0)))
    row_admin_total = total('ИТОГО АДМИНИСТРАТИВНЫЕ', rowrefs=adm_rows, fill='BDD7EE')
    r += 1

    # ---------------- ПРОЧИЕ ДОХОДЫ И РАСХОДЫ (НИЗ) ----------------
    section('ПРОЧИЕ ДОХОДЫ И РАСХОДЫ')
    r_dep = leaf('Доходы от финансовой деятельности (депозиты)', lambda R: R['deposit'])
    r_fxi = leaf('Доход от курсовой разницы', lambda R: R['fx_income'])
    r_oi  = leaf('Прочие доходы', lambda R: R['other_income'])
    r_fxe = leaf('Расход по курсовой разнице', lambda R: R['fx_expense'])
    r_fin = leaf('Расходы на финансирование', lambda R: R['financing'])
    r_oe  = leaf('Прочие расходы', lambda R: R['other_expense'])
    r += 1

    # ---------------- ИТОГИ ----------------
    row_pbt = total('ДОХОД ДО КПН',
        formula=lambda col: (f'={col}{row_gross}-{col}{row_admin_total}'
                             f'+{col}{r_dep}+{col}{r_fxi}+{col}{r_oi}'
                             f'-{col}{r_fxe}-{col}{r_fin}-{col}{r_oe}'),
        fill='FCE4D6')
    row_kpn = total('КПН 20%',
        formula=lambda col: f'=MAX(0,{col}{row_pbt})*0.2', fill='FCE4D6')
    row_net = total('ЧИСТЫЙ ДОХОД',
        formula=lambda col: f'={col}{row_pbt}-{col}{row_kpn}', fill='C6E0B4')
    r += 1

    # ---------------- КОНТРОЛЬ СВЕРКИ С ОСВ ----------------
    section('КОНТРОЛЬ СВЕРКИ С ОСВ 1С')
    r_i6 = leaf('Обороты 6ххх (доходы, Кт)', lambda R: R['tot_income_6'])
    r_e7 = leaf('Обороты 7ххх (расходы, Дт)', lambda R: R['tot_expense_7'])
    r_res = total('Результат 1С (6ххх − 7ххх)',
        formula=lambda col: f'={col}{r_i6}-{col}{r_e7}', fill=None, bold=False)
    r_diff = total('Разница (P&L − 1С)',
        formula=lambda col: f'={col}{row_pbt}-{col}{r_res}', fill='FFFF00')

    # ширины колонок
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 52
    for col in data_cols + [tot_col]:
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = 'C2'

    wb.save(out_path)
    return {'rows': r, 'projects': len(projects), 'pbt_row': row_pbt}


def build_journal(out_path, periods):
    """Добавляет лист «Журнал сопоставления»: какие сырые проекты 1С свёрнуты в блок,
    группировка налогов, суммы по статьям и контроль."""
    wb = openpyxl.load_workbook(out_path)
    if 'Журнал сопоставления' in wb.sheetnames:
        del wb['Журнал сопоставления']
    js = wb.create_sheet('Журнал сопоставления')
    hdr = ['Период', 'Раздел', 'Блок/Статья в P&L', 'Исходное имя из 1С', 'Сумма ₸']
    js.append(hdr)
    for c in js[1]:
        c.font = Font(name=_F_NAME, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F5496'); c.border = _BORDER
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    for nm, R in periods:
        # проекты: сырые имена -> блок
        for proj in sorted(R['proj_raw']):
            raws = sorted(R['proj_raw'][proj])
            rev = R['income'].get(proj, 0.0)
            cost = sum(R['cost'].get(proj, {}).values())
            for raw in raws:
                if norm(raw) != norm(proj):
                    js.append([nm, 'Проект (свёрнут)', proj, raw, ''])
            js.append([nm, 'Проект — выручка', proj, '', round(rev, 2)])
            js.append([nm, 'Проект — затраты', proj, '', round(cost, 2)])
        # налоги с ЗП — что вошло
        tax_parts = {a: v for a, v in R['admin'].items()
                     if any(k in norm(a) for k in _TAX_KEYS)}
        for a, v in tax_parts.items():
            js.append([nm, 'Налоги по ЗП (свёрнуто)', 'Налоги и обяз. платежи по ЗП', a, round(v, 2)])
        # низ
        for lbl, key in [('Доходы от фин. деятельности', 'deposit'),
                         ('Доход от курсовой разницы', 'fx_income'),
                         ('Прочие доходы', 'other_income'),
                         ('Расход по курсовой разнице', 'fx_expense'),
                         ('Расходы на финансирование', 'financing'),
                         ('Прочие расходы', 'other_expense')]:
            if R.get(key):
                js.append([nm, 'Прочие доходы/расходы', lbl, '', round(R[key], 2)])
        # контроль
        js.append([nm, 'КОНТРОЛЬ', 'Обороты 6ххх', '', round(R['tot_income_6'], 2)])
        js.append([nm, 'КОНТРОЛЬ', 'Обороты 7ххх', '', round(R['tot_expense_7'], 2)])
        js.append([nm, 'КОНТРОЛЬ', 'Результат 1С (6−7)', '', round(R['tot_income_6'] - R['tot_expense_7'], 2)])
    for i, w in enumerate([12, 24, 40, 46, 18], 1):
        js.column_dimensions[get_column_letter(i)].width = w
    for row in js.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=_F_NAME, size=9); c.border = _BORDER
        row[4].number_format = '#,##0'
    js.freeze_panes = 'A2'
    wb.save(out_path)
    return js.max_row - 1
